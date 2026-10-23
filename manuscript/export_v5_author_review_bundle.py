#!/usr/bin/env python3
"""Generate non-final Azores V5 Word and supplementary PDF using unchanged V4 science.

Pandoc/LibreOffice file conversion only. Source V4 numeric contract,
original V5 manuscript, 35 verbatim S1–S8 supplement passages and verified
four SVG figures are never recalculated or edited. Authors remain placeholders.
"""
from __future__ import annotations
import hashlib,json,subprocess,shutil,zipfile
from pathlib import Path

import cairosvg, fitz, pypandoc
from docx import Document
from docx.shared import Pt,Inches,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
MAIN=ROOT/"manuscript/AZORES_PHASE_TRANSFER_MAIN_V5_DRAFT.md"
SUPP=ROOT/"manuscript/AZORES_V5_SUPPLEMENT_METHODS_RESULTS_DRAFT.md"
V4=ROOT/"manuscript/AZORES_PHASE_CONTROL_MANUSCRIPT_V4.md"
CONTRACT=ROOT/"manuscript/FIGURE_DATA_CONTRACT_V4.json"
FIGDIR=ROOT/"manuscript/rendered_figures_v4"
OUT=ROOT/"results/generated_v5_submission_review"
ARCHIVE=ROOT/"results/AZORES_V5_AUTHOR_REVIEW_PACKAGE.zip"

def hash_file(f):
    return hashlib.sha256(f.read_bytes()).hexdigest()

def style_doc(p):
    d=Document(str(p))
    for name in ("Normal","Body Text","Heading 1","Heading 2","Heading 3"):
        if name not in d.styles:continue
        s=d.styles[name]
        s.font.name="Times New Roman"
        s.font.color.rgb=RGBColor(28,43,51)
        s.font.size=Pt({"Heading 1":12,"Heading 2":11,"Heading 3":10}.get(name,10))
        if name in ("Normal","Body Text"):
            s.paragraph_format.line_spacing=1.4
            s.paragraph_format.space_after=Pt(5)
        else:
            s.font.bold=True
            s.paragraph_format.space_before=Pt(10)
            s.paragraph_format.space_after=Pt(5)
            s.paragraph_format.keep_with_next=True
    for section in d.sections:
        section.top_margin=Inches(.85)
        section.bottom_margin=Inches(.8)
        section.left_margin=Inches(.89)
        section.right_margin=Inches(.87)
        section.header_distance=Inches(.4)
        section.footer_distance=Inches(.42)
        p=section.header.paragraphs[0]
        p.text="AUTHOR REVIEW  |  SCIENTIFIC DRAFT — NOT SUBMITTED"
        p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        for run in p.runs:run.font.size=Pt(7)
        foot=section.footer.paragraphs[0]
        foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        foot.add_run("Page ").font.size=Pt(8)
        field=OxmlElement("w:fldSimple")
        field.set(qn("w:instr"),"PAGE")
        foot._p.append(field)
        l=OxmlElement("w:lnNumType")
        for k,v in [("countBy","1"),("start","1"),("distance","360"),("restart","continuous")]:
            l.set(qn("w:"+k),v)
        for old in section._sectPr.findall(qn("w:lnNumType")):section._sectPr.remove(old)
        section._sectPr.append(l)
    d.core_properties.author="Actual contributors not yet supplied"
    d.core_properties.subject="Azores V5 public-source ecological analysis, author-review text export only"
    d.save(str(p))

def docx(md,out,resources):
    src=out.with_suffix(".input.md")
    src.write_text(md,encoding="utf-8")
    args=["--standalone","--resource-path="+":".join(str(x) for x in resources)]
    pypandoc.convert_file(str(src),"docx",format="markdown-yaml_metadata_block+tex_math_dollars",
                          outputfile=str(out),extra_args=args)
    style_doc(out)
    src.unlink()
    return out

def pdf_from(doc,out):
    made=OUT/"lo_pdf"
    made.mkdir(exist_ok=True)
    p=subprocess.run(["libreoffice","-env:UserInstallation=file:///tmp/azores_v5_authorproof",
                      "--headless","--convert-to","pdf","--outdir",str(made),str(doc)],
                     text=True,capture_output=True,timeout=220)
    dest=made/(doc.stem+".pdf")
    if p.returncode or not dest.is_file():
        raise RuntimeError("LibreOffice failed: "+p.stdout[:350]+" "+p.stderr[:350])
    shutil.copy2(dest,out)
    return out

def audit(pdf,required,minimum_pages,figures=0):
    d=fitz.open(str(pdf))
    assert d.page_count>=minimum_pages,(pdf.name,d.page_count)
    text="\n".join(p.get_text(sort=True) for p in d)
    for needle in required:
        assert needle.lower() in text.lower(),(pdf.name,needle)
    count=0
    for page in d:
        count+=len(page.get_images(full=True))
        for b in page.get_text("blocks"):
            if len(b)<5 or not isinstance(b[4],str):continue
            x0,y0,x1,y1=b[:4]
            assert x0>=-2 and y0>=-2 and x1<=page.rect.width+2 and y1<=page.rect.height+2, (pdf.name,page.number,b)
    assert count>=figures,(pdf.name,count)
    return {"pages":d.page_count,"rendered_image_references":count,"bbox_clipping":0}

def main():
    m=MAIN.read_text(encoding="utf-8")
    s=SUPP.read_text(encoding="utf-8")
    assert "2.23" in m and "0.946" in m and "Albertkanaal" in m
    assert s.count("<!-- V4_SOURCE_BEGIN:")==35
    assert "## S8." in s and "## S1." in s
    OUT.mkdir(parents=True,exist_ok=True)
    prefix="**NON-FINAL AUTHOR REVIEW PROOF — VERIFY ALL CONTRIBUTORS, ETHICS, REFERENCES AND SOURCE RIGHTS BEFORE SUBMISSION.**\n\n"
    paper=docx(prefix+m,OUT/"Azores_V5_Main_AuthorReview.docx",[MAIN.parent])
    supp=docx(prefix+s,OUT/"Azores_V5_Supplement_AuthorReview.docx",[SUPP.parent])
    pngs=[]
    figures=OUT/"source_figures"
    figures.mkdir(exist_ok=True)
    for i in range(1,5):
        src=FIGDIR/f"Figure{i}.svg"
        assert src.is_file() and src.stat().st_size>1000
        png=figures/f"Figure{i}_proof.png"
        cairosvg.svg2png(url=str(src),write_to=str(png),scale=1.5)
        assert png.stat().st_size>1000
        shutil.copy2(src,figures/src.name)
        pngs.append(png)
    figure_md=(
       "# Original Azores V4 figures — Author review proof\n\n"
       "**These are the four existing source-backed figures. This PDF contains no new analysis.**\n\n"
       "Inspect figure lettering and endpoint interpretation before journal upload.\n\n"
    )
    for i in range(1,5):
        figure_md+=f"## Figure {i}\n\n![V4 source Figure {i}](Figure{i}_proof.png){{ width=6.2in }}\n\n"
    figure_doc=docx(figure_md,OUT/"Azores_V5_Figures_AuthorReview.docx",[figures])
    ppdf=pdf_from(paper,OUT/"Azores_V5_Main_PROOF.pdf")
    spdf=pdf_from(supp,OUT/"Azores_V5_Supplement_AuthorReview.pdf")
    fpdf=pdf_from(figure_doc,OUT/"Azores_V5_FigureBook_AuthorReview.pdf")
    qc={
      "status":"AUTHOR_REVIEW_EXPORT_NOT_SUBMITTED",
      "source_sha256":{str(x.relative_to(ROOT)):hash_file(x) for x in [MAIN,SUPP,V4,CONTRACT]},
      "figures":4,"verbatim_supplement_blocks":35,
      "pdf_checks":{
        "main":audit(ppdf,["Migration activation","Abstract","Methods","Discussion","Albertkanaal"],7),
        "supplement":audit(spdf,["Supplementary Methods and Results","S8.","FIII","Albertkanaal"],10),
        "figures":audit(fpdf,["Figure 1","Figure 4","review"],3,4)},
      "visual_review_required":True,
      "author_submission_blockers":[
        "actual author list/affiliations/corresponding email","study ethics/data permission verification",
        "full publisher-verified bibliography","journal-specific supplement tables/figures and Word layout proof",
        "submission DOI/repository release and author approvals"],
    }
    (OUT/"AUTHOR_CHECKLIST.md").write_text(
      "# Azores V5 scientific paper author review\n\nNo manuscript was submitted.\n\n"+
      "\n".join("- [ ] "+v for v in qc["author_submission_blockers"])+"\n",
      encoding="utf-8")
    (OUT/"SOURCE_AND_LAYOUT_QC.json").write_text(json.dumps(qc,indent=2)+"\n",encoding="utf-8")
    shutil.rmtree(OUT/"lo_pdf",ignore_errors=True)
    ARCHIVE.parent.mkdir(parents=True,exist_ok=True)
    with zipfile.ZipFile(ARCHIVE,"w",compression=zipfile.ZIP_DEFLATED,compresslevel=6) as z:
        for f in sorted(OUT.rglob("*")):
            if f.is_file():z.write(f,f.relative_to(OUT))
    print(json.dumps({"status":qc["status"],"archive":str(ARCHIVE),
          "bytes":ARCHIVE.stat().st_size,"sha256":hash_file(ARCHIVE),
          "pdf_checks":qc["pdf_checks"]},indent=2))

if __name__=="__main__":
    main()
